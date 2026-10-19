"""Offline Excel routing tests: temporary workbooks only, never WeChat or Excel UI."""

from pathlib import Path
import subprocess
import sys
import tempfile
import time
import unittest
from unittest.mock import patch
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.table import Table, TableStyleInfo

from wechat import group_mapping as mapping


class GroupMappingTests(unittest.TestCase):
    def setUp(self):
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.root = Path(directory.name)
        self.path = self.root / "data" / mapping.DEFAULT_MAPPING_FILENAME
        self.store = mapping.GroupMappingStore(self.path)

    def write_rows(self, rows, *, headers=None):
        self.path.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        workbook.active.append(headers or list(mapping.HEADERS))
        for row in rows:
            workbook.active.append(row)
        workbook.save(self.path)
        workbook.close()

    def edit(self, callback):
        workbook = load_workbook(self.path)
        callback(workbook)
        workbook.save(self.path)
        workbook.close()

    def test_create_empty_template_and_resolve_requires_explicit_mapping(self):
        self.assertEqual(self.store.create(), self.path)
        self.assertEqual(self.store.load(), [])
        workbook = load_workbook(self.path)
        self.assertEqual(tuple(cell.value for cell in workbook.active[1]), mapping.HEADERS)
        self.assertEqual(workbook.active.freeze_panes, "A2")
        workbook.close()
        with self.assertRaisesRegex(mapping.MappingError, "没有公司.*不会使用公司名搜索"):
            self.store.resolve("甲公司")

    def test_missing_workbook_is_not_created_by_resolve_or_upsert(self):
        for operation in (lambda: self.store.resolve("甲公司"), lambda: self.store.upsert("甲公司", "财税服务群")):
            with self.subTest(operation=operation), self.assertRaisesRegex(mapping.MappingError, "新建 Excel"):
                operation()
            self.assertFalse(self.path.exists())

    def test_manual_two_column_workbook_can_add_mapping(self):
        self.write_rows([], headers=list(mapping.HEADERS[:2]))
        entry = self.store.upsert("甲公司", "甲司财税服务群")
        self.assertEqual(entry.company_name, "甲公司")
        self.assertEqual(entry.group_name, "甲司财税服务群")
        self.assertTrue(entry.updated_at)
        self.assertEqual(self.store.resolve("甲公司"), "甲司财税服务群")

    def test_lookup_is_exact_with_no_substring_or_normalization(self):
        self.write_rows([("甲 公司（上海）", "ABC 甲司群 🧾")])
        self.assertEqual(self.store.resolve("甲 公司（上海）"), "ABC 甲司群 🧾")
        for company in ("甲", "甲公司（上海）", "甲 公司(上海)"):
            with self.subTest(company=company), self.assertRaises(mapping.MappingError):
                self.store.resolve(company)

    def test_preserve_leading_zeros_and_unicode_group_name(self):
        self.write_rows([("001", "001 财税群（正式）🧾")])
        self.assertEqual(self.store.resolve("001"), "001 财税群（正式）🧾")

    def test_reloads_changes_made_in_excel_instead_of_caching(self):
        self.write_rows([("甲公司", "旧群")])
        self.assertEqual(self.store.resolve("甲公司"), "旧群")
        self.edit(lambda book: setattr(book.active["B2"], "value", "新群"))
        self.assertEqual(self.store.resolve("甲公司"), "新群")

    def test_upsert_updates_company_without_creating_duplicates(self):
        self.write_rows([("甲公司", "旧群", "保留备注"), ("乙公司", "乙群")])
        result = self.store.upsert("甲公司", "新群")
        self.assertEqual(result.notes, "保留备注")
        self.assertEqual([entry.group_name for entry in self.store.load()], ["新群", "乙群"])

    def test_multiple_companies_may_share_same_group(self):
        self.write_rows([("甲公司", "集团财税群")])
        self.store.upsert("乙公司", "集团财税群")
        self.assertEqual(self.store.resolve("乙公司"), "集团财税群")
        self.assertEqual(len(self.store.load()), 2)

    def test_expected_absent_mapping_can_be_created_after_verification(self):
        self.store.create()
        self.store.upsert("甲公司", "新建微信群", expected_group=None)
        self.assertEqual(self.store.resolve("甲公司"), "新建微信群")

    def test_expected_original_group_can_be_updated_after_verification(self):
        self.write_rows([("甲公司", "原群")])
        self.store.upsert("甲公司", "已核验新群", expected_group="原群")
        self.assertEqual(self.store.resolve("甲公司"), "已核验新群")

    def test_edit_while_waiting_does_not_overwrite_new_group(self):
        self.write_rows([("甲公司", "开始等待时的群")])
        expected = self.store.resolve("甲公司")
        self.edit(lambda book: setattr(book.active["B2"], "value", "等待期间另行修正的群"))
        original = self.path.read_bytes()
        with self.assertRaisesRegex(mapping.MappingError, "等待或验证期间已发生变化.*刷新"):
            self.store.upsert("甲公司", "刚建好的群", expected_group=expected)
        self.assertEqual(self.path.read_bytes(), original)
        self.assertEqual(self.store.resolve("甲公司"), "等待期间另行修正的群")

    def test_new_company_added_while_waiting_is_not_overwritten(self):
        self.store.create()
        self.edit(lambda book: book.active.append(("甲公司", "他人登记的群")))
        original = self.path.read_bytes()
        with self.assertRaisesRegex(mapping.MappingError, "现有配对未覆盖"):
            self.store.upsert("甲公司", "刚建好的群", expected_group=None)
        self.assertEqual(self.path.read_bytes(), original)

    def test_company_removed_while_waiting_is_not_recreated(self):
        self.write_rows([("甲公司", "原群")])
        self.edit(lambda book: book.active.delete_rows(2))
        original = self.path.read_bytes()
        with self.assertRaisesRegex(mapping.MappingError, "现有配对未覆盖"):
            self.store.upsert("甲公司", "刚建好的群", expected_group="原群")
        self.assertEqual(self.path.read_bytes(), original)

    def test_unrelated_edits_while_waiting_are_preserved_and_do_not_block(self):
        self.write_rows([("甲公司", "原群")])
        self.edit(lambda book: book.active.append(("乙公司", "乙群", "新备注")))
        self.store.upsert("甲公司", "刚建好的群", expected_group="原群")
        self.assertEqual(self.store.resolve("甲公司"), "刚建好的群")
        self.assertEqual(self.store.resolve("乙公司"), "乙群")
        self.assertEqual(self.store.list_entries()[1].notes, "新备注")

    def test_duplicate_company_rejected_even_if_group_names_match(self):
        for second in ("原群", "另一个群"):
            self.write_rows([("甲公司", "原群"), ("甲公司", second)])
            original = self.path.read_bytes()
            with self.subTest(second=second):
                for operation in (lambda: self.store.resolve("甲公司"), lambda: self.store.upsert("乙公司", "乙群")):
                    with self.assertRaisesRegex(mapping.MappingError, "第 2、3 行重复"):
                        operation()
                self.assertEqual(self.path.read_bytes(), original)

    def test_invalid_required_cells_rejected_before_any_destination_is_resolved(self):
        invalid = (None, 123, True, "", " ", "=1+1", "=\"财税群\"", " 群", "群 ", "群\n名称")
        for value in invalid:
            for column in (0, 1):
                row = ["错误公司", "错误群"]
                row[column] = value
                self.write_rows([("正常公司", "正常群"), row])
                with self.subTest(value=value, column=column), self.assertRaises(mapping.MappingError):
                    self.store.resolve("正常公司")

    def test_blank_rows_ignored_but_partial_rows_rejected(self):
        self.write_rows([(None, None, "单独备注"), ("甲公司", "甲群")])
        self.assertEqual(len(self.store.load()), 1)
        self.store.upsert("乙公司", "乙群")
        book = load_workbook(self.path)
        self.assertEqual(book.active["C2"].value, "单独备注")
        self.assertEqual(book.active["A4"].value, "乙公司")
        book.close()

    def test_formatted_blank_rows_are_reused_instead_of_appending_far_below(self):
        self.write_rows([("甲公司", "甲群")])
        def style_blank_rows(book):
            for row in range(3, 501):
                book.active.cell(row, 1).font = Font(bold=True)
        self.edit(style_blank_rows)
        self.store.upsert("新公司", "新群")
        book = load_workbook(self.path)
        self.assertEqual(book.active["A3"].value, "新公司")
        self.assertEqual(book.active["B3"].value, "新群")
        self.assertTrue(book.active["A3"].font.bold)
        self.assertEqual(book.active.max_row, 500)
        book.close()

    def test_reused_blank_row_never_overwrites_notes_or_extra_column_values(self):
        self.write_rows([
            ("甲公司", "甲群"),
            (None, None, "不要覆盖的备注"),
            (None, None, None, None, "额外列内容"),
        ], headers=[*mapping.HEADERS, "负责人"])
        self.edit(lambda book: setattr(book.active["A200"], "font", Font(bold=True)))
        self.store.upsert("新公司", "新群")
        book = load_workbook(self.path)
        self.assertEqual(book.active["C3"].value, "不要覆盖的备注")
        self.assertEqual(book.active["E4"].value, "额外列内容")
        self.assertIsNone(book.active["A3"].value)
        self.assertIsNone(book.active["A4"].value)
        self.assertEqual(book.active["A5"].value, "新公司")
        book.close()

    def test_existing_company_updates_original_row_despite_earlier_blank_rows(self):
        self.write_rows([])
        def populate_later_row(book):
            book.active["A20"] = "已有公司"
            book.active["B20"] = "原群"
        self.edit(populate_later_row)
        self.store.upsert("已有公司", "更新后的群")
        book = load_workbook(self.path)
        self.assertIsNone(book.active["A2"].value)
        self.assertEqual(book.active["A20"].value, "已有公司")
        self.assertEqual(book.active["B20"].value, "更新后的群")
        book.close()

    def test_invalid_input_never_changes_workbook(self):
        self.store.create()
        original = self.path.read_bytes()
        for value in (None, 1, True, "", " ", "=SUM(A1)", " 名称", "名称 ", "名\t称", "名\x00称", "x" * 32768):
            for company, group in ((value, "微信群"), ("公司", value)):
                with self.subTest(value=repr(value)[:25]), self.assertRaises(mapping.MappingError):
                    self.store.upsert(company, group)
                self.assertEqual(self.path.read_bytes(), original)

    def test_notes_can_be_preserved_updated_or_cleared(self):
        self.write_rows([("甲公司", "甲群", "原备注")])
        self.assertEqual(self.store.upsert("甲公司", "新群").notes, "原备注")
        self.assertEqual(self.store.upsert("甲公司", "新群", notes="新备注").notes, "新备注")
        self.assertEqual(self.store.upsert("甲公司", "新群", notes="").notes, "")

    def test_formula_like_notes_are_saved_as_text(self):
        self.write_rows([], headers=list(mapping.HEADERS[:2]))
        self.store.upsert("甲公司", "甲群", notes="=HYPERLINK(\"https://example.test\")")
        book = load_workbook(self.path, data_only=False)
        self.assertEqual(book.active["C2"].data_type, "s")
        self.assertTrue(book.active["C2"].value.startswith("="))
        book.close()

    def test_last_formatted_column_does_not_create_invalid_excel_columns(self):
        self.write_rows([], headers=list(mapping.HEADERS[:2]))
        self.edit(lambda book: setattr(book.active["XFD1"], "font", Font(bold=True)))
        self.store.upsert("甲公司", "甲群", notes="新备注")
        book = load_workbook(self.path)
        self.assertEqual(book.active.max_column, 16384)
        self.assertEqual(book.active["C1"].value, "备注")
        self.assertEqual(book.active["C2"].value, "新备注")
        self.assertEqual(book.active["D1"].value, "更新时间")
        self.assertTrue(book.active["XFD1"].font.bold)
        book.close()

    def test_far_corner_formatting_keeps_reads_and_new_rows_sparse(self):
        self.write_rows([("甲公司", "甲群")], headers=list(mapping.HEADERS[:2]))
        self.edit(lambda book: setattr(book.active["XFD1048576"], "font", Font(bold=True)))
        book, sheet, _original = self.store._open()
        try:
            before = len(sheet._cells)
            self.assertEqual([entry.company_name for _, entry in self.store._entries(sheet)], ["甲公司"])
            self.assertEqual(len(sheet._cells), before)
            self.assertEqual(self.store._first_empty_row(sheet), 3)
            self.assertEqual(len(sheet._cells), before)
        finally:
            book.close()
        self.store.upsert("乙公司", "乙群", notes="新登记")
        book = load_workbook(self.path)
        self.assertEqual(book.active["A3"].value, "乙公司")
        self.assertEqual(book.active["B3"].value, "乙群")
        self.assertTrue(book.active["XFD1048576"].font.bold)
        self.assertEqual(book.active.max_column, 16384)
        self.assertEqual(book.active.max_row, 1048576)
        self.assertLess(len(book.active._cells), 20)
        book.close()

    def test_new_optional_column_preserves_data_below_blank_header(self):
        self.write_rows([("甲公司", "甲群", "未命名列中的资料")], headers=list(mapping.HEADERS[:2]))
        self.store.upsert("甲公司", "新群")
        book = load_workbook(self.path)
        self.assertIsNone(book.active["C1"].value)
        self.assertEqual(book.active["C2"].value, "未命名列中的资料")
        self.assertEqual(book.active["D1"].value, "更新时间")
        book.close()

    def test_no_optional_column_space_preserves_original_workbook(self):
        self.write_rows([], headers=[*mapping.HEADERS[:2], *[f"占用列 {n}" for n in range(3, 16385)]])
        original = self.path.read_bytes()
        with self.assertRaisesRegex(mapping.MappingError, "没有可用的空白列"):
            self.store.upsert("甲公司", "甲群")
        self.assertEqual(self.path.read_bytes(), original)

    def test_extra_columns_formatting_formulas_and_other_sheets_survive(self):
        self.write_rows([("甲公司", "甲群", "批注", "负责人甲", "=1+1")], headers=[*mapping.HEADERS[:3], "负责人", "核算"])
        def setup(book):
            book.active["D2"].font = Font(bold=True, color="FF0000")
            book.create_sheet("其他数据")["A1"] = "保留其他数据"
        self.edit(setup)
        self.store.upsert("甲公司", "新群")
        book = load_workbook(self.path, data_only=False)
        self.assertEqual(book.active["C2"].value, "批注")
        self.assertEqual(book.active["D2"].value, "负责人甲")
        self.assertTrue(book.active["D2"].font.bold)
        self.assertEqual(book.active["E2"].value, "=1+1")
        self.assertEqual(book["其他数据"]["A1"].value, "保留其他数据")
        book.close()

    def test_unknown_layout_never_overwritten_even_by_create(self):
        self.write_rows([("甲公司", "甲群")], headers=["公司", "群名"])
        original = self.path.read_bytes()
        for operation in (self.store.create, lambda: self.store.upsert("甲公司", "新群")):
            with self.assertRaisesRegex(mapping.MappingError, "A1"):
                operation()
            self.assertEqual(self.path.read_bytes(), original)

    def test_create_existing_valid_file_is_nonmutating(self):
        self.write_rows([("甲公司", "甲群")])
        original = self.path.read_bytes()
        self.store.create()
        self.assertEqual(self.path.read_bytes(), original)

    def test_multiple_matching_sheets_rejected(self):
        self.store.create()
        self.edit(lambda book: book.create_sheet("另一张配对表").append(mapping.HEADERS))
        with self.assertRaisesRegex(mapping.MappingError, "多张配对表"):
            self.store.load()

    def test_duplicate_optional_header_rejected(self):
        self.write_rows([("甲公司", "甲群")], headers=[*mapping.HEADERS, "备注"])
        with self.assertRaisesRegex(mapping.MappingError, "表头.*重复"):
            self.store.load()

    def test_merged_mapping_cells_rejected(self):
        self.write_rows([("甲公司", "甲群")])
        self.edit(lambda book: book.active.merge_cells("A2:A3"))
        with self.assertRaisesRegex(mapping.MappingError, "合并单元格"):
            self.store.load()

    def test_protected_sheet_not_modified(self):
        self.write_rows([("甲公司", "甲群")])
        self.edit(lambda book: setattr(book.active.protection, "sheet", True))
        original = self.path.read_bytes()
        with self.assertRaisesRegex(mapping.MappingError, "受保护"):
            self.store.upsert("甲公司", "新群")
        self.assertEqual(self.path.read_bytes(), original)

    def test_excel_owner_lock_blocks_writes_with_actionable_message(self):
        self.store.create()
        original = self.path.read_bytes()
        self.path.with_name("~$" + self.path.name).write_bytes(b"Excel owner file")
        with self.assertRaisesRegex(mapping.MappingError, "保存并关闭 Excel"):
            self.store.upsert("甲公司", "甲群")
        self.assertEqual(self.path.read_bytes(), original)

    def test_replace_permission_error_preserves_original_and_removes_temp_file(self):
        self.write_rows([("甲公司", "甲群")])
        original = self.path.read_bytes()
        with patch.object(mapping.os, "replace", side_effect=PermissionError("locked")):
            with self.assertRaisesRegex(mapping.MappingError, "关闭 Excel"):
                self.store.upsert("甲公司", "新群")
        self.assertEqual(self.path.read_bytes(), original)
        self.assertEqual(list(self.path.parent.iterdir()), [self.path])

    def test_save_failure_preserves_original_and_removes_temp_file(self):
        self.write_rows([("甲公司", "甲群")])
        original = self.path.read_bytes()
        with patch.object(Workbook, "save", side_effect=OSError("disk full")):
            with self.assertRaises(mapping.MappingError):
                self.store.upsert("甲公司", "新群")
        self.assertEqual(self.path.read_bytes(), original)
        self.assertEqual(list(self.path.parent.iterdir()), [self.path])

    def test_external_edit_during_save_is_not_overwritten(self):
        self.write_rows([("甲公司", "甲群")])
        actual_save = Workbook.save
        external = b"new external contents sentinel"
        def racing_save(book, filename):
            actual_save(book, filename)
            self.path.write_bytes(external)
        with patch.object(Workbook, "save", racing_save):
            with self.assertRaisesRegex(mapping.MappingError, "其他程序修改"):
                self.store.upsert("甲公司", "新群")
        self.assertEqual(self.path.read_bytes(), external)

    def start_writer_paused_before_replace(self):
        """A real second process holds the write transaction at its last step."""
        ready, release = self.root / "writer-ready", self.root / "writer-release"
        script = '''from pathlib import Path
import sys, time
from wechat import group_mapping as mapping
path, ready, release = map(Path, sys.argv[1:])
replace = mapping.os.replace
def paused_replace(source, destination):
    ready.touch()
    deadline = time.monotonic() + 20
    while not release.exists():
        if time.monotonic() > deadline:
            raise RuntimeError("test writer was not released")
        time.sleep(0.01)
    return replace(source, destination)
mapping.os.replace = paused_replace
mapping.GroupMappingStore(path).upsert("甲公司", "甲群")
'''
        process = subprocess.Popen(
            [sys.executable, "-c", script, str(self.path), str(ready), str(release)],
            cwd=Path(mapping.__file__).resolve().parents[1],
            stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True,
        )
        def stop_writer():
            if process.poll() is None:
                process.kill()
            process.communicate(timeout=10)
        self.addCleanup(stop_writer)
        deadline = time.monotonic() + 15
        while not ready.exists():
            if process.poll() is not None or time.monotonic() > deadline:
                stop_writer()
                self.fail("second process did not reach its workbook commit")
            time.sleep(0.01)
        return process, release

    def test_concurrent_process_cannot_overwrite_successful_mapping(self):
        self.store.create()
        original = self.path.read_bytes()
        process, release = self.start_writer_paused_before_replace()
        # Equivalent path spellings must share a lock, and a different workbook
        # must remain writable while this one has a pending transaction.
        alias = self.path.parent / ".." / self.path.parent.name / self.path.name
        with self.assertRaisesRegex(mapping.MappingError, "另一个任务正在保存"):
            mapping.GroupMappingStore(alias).upsert("乙公司", "乙群")
        self.assertEqual(self.path.read_bytes(), original)
        independent = mapping.GroupMappingStore(self.root / "另一张配对表.xlsx")
        independent.create()
        independent.upsert("丙公司", "丙群")
        release.touch()
        output, errors = process.communicate(timeout=15)
        self.assertEqual(process.returncode, 0, output + errors)
        self.store.upsert("乙公司", "乙群")
        self.assertEqual(
            [(entry.company_name, entry.group_name) for entry in self.store.load()],
            [("甲公司", "甲群"), ("乙公司", "乙群")],
        )

    def test_process_exit_releases_write_lock_without_manual_cleanup(self):
        self.store.create()
        process, _release = self.start_writer_paused_before_replace()
        process.kill()
        process.communicate(timeout=15)
        # The interrupted write never replaced the workbook, and its OS lock
        # cannot strand future registrations behind a stale lock file.
        self.assertEqual(self.store.load(), [])
        self.store.upsert("乙公司", "乙群")
        self.assertEqual(self.store.resolve("乙公司"), "乙群")

    def test_external_creation_is_not_overwritten_by_template(self):
        actual_link = mapping.os.link
        external = b"another process created this file"
        def racing_link(source, destination):
            self.path.write_bytes(external)
            return actual_link(source, destination)
        with patch.object(mapping.os, "link", racing_link):
            with self.assertRaisesRegex(mapping.MappingError, "其他程序创建"):
                self.store.create()
        self.assertEqual(self.path.read_bytes(), external)

    def test_unsupported_excel_content_is_not_silently_removed(self):
        self.store.create()
        with ZipFile(self.path, "a") as archive:
            archive.writestr("xl/drawings/drawing1.xml", b"test unsupported content")
        original = self.path.read_bytes()
        with self.assertRaisesRegex(mapping.MappingError, "高级内容"):
            self.store.upsert("甲公司", "甲群")
        self.assertEqual(self.path.read_bytes(), original)

    def test_corrupt_file_and_unsupported_extensions_have_clear_errors(self):
        self.path.parent.mkdir(parents=True)
        self.path.write_bytes(b"not an Excel workbook")
        with self.assertRaisesRegex(mapping.MappingError, "无法打开"):
            self.store.load()
        for suffix in (".xls", ".xlsm", ".csv"):
            with self.subTest(suffix=suffix), self.assertRaisesRegex(mapping.MappingError, "xlsx"):
                mapping.GroupMappingStore(self.path.with_suffix(suffix))

    def test_mapping_table_filter_expands_to_added_record(self):
        self.write_rows([("甲公司", "甲群")], headers=list(mapping.HEADERS[:2]))
        def setup(book):
            table = Table(displayName="Mapping", ref="A1:B2")
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
            book.active.add_table(table)
            book.active.auto_filter.ref = "A1:B2"
        self.edit(setup)
        self.store.upsert("乙公司", "乙群")
        book = load_workbook(self.path)
        self.assertEqual(book.active.tables["Mapping"].ref, "A1:B3")
        self.assertEqual(book.active.auto_filter.ref, "A1:B3")
        book.close()


if __name__ == "__main__":
    unittest.main()
