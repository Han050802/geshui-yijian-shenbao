"""Exact company-to-WeChat-group mappings stored in an ordinary Excel workbook.

This module only reads and writes local files. It does not inspect the clipboard,
connect to WeChat, or infer a destination when an Excel entry is missing.
"""

from __future__ import annotations

from contextlib import contextmanager
from dataclasses import dataclass
from datetime import datetime
from hashlib import sha256
from io import BytesIO
import os
from pathlib import Path
import tempfile
import threading
import warnings
from zipfile import BadZipFile, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.cell import get_column_letter, range_boundaries


DEFAULT_MAPPING_FILENAME = "微信群配对.xlsx"
SHEET_NAME = "微信群配对"
HEADERS = ("公司名称", "微信群完整名称", "备注", "更新时间")
_WRITE_LOCK = threading.RLock()
_UNSPECIFIED = object()


class MappingError(ValueError):
    """An actionable validation or file-access error in the mapping workbook."""


@dataclass(frozen=True)
class GroupMappingEntry:
    company_name: str
    group_name: str
    notes: str = ""
    updated_at: str = ""


def _name(value: object, label: str) -> str:
    if not isinstance(value, str) or not value or not value.strip():
        raise MappingError(f"{label}必须填写非空文本，不能使用数字、公式或空白。")
    if value != value.strip():
        raise MappingError(f"{label}首尾有空格，请在 Excel 和微信中核对并去除后重试。")
    if value.startswith("="):
        raise MappingError(f"{label}不能使用公式，请填写完整名称的纯文本。")
    if len(value) > 32767 or any(ord(char) < 32 for char in value):
        raise MappingError(f"{label}包含换行、控制字符或超长文本，请填写单行完整名称。")
    return value


def _optional_text(value: object, label: str) -> str:
    if value is None:
        return ""
    if not isinstance(value, str):
        return str(value)
    return value


class GroupMappingStore:
    """Read a fresh workbook for every operation; never cache a routing decision.

    A manually created .xlsx file is supported when exactly one worksheet has
    A1=公司名称 and B1=微信群完整名称. Other columns and sheets are retained.
    Each company may appear once; several companies may share the same group.
    """

    def __init__(self, path: str | Path) -> None:
        self.path = Path(path)
        if self.path.suffix.lower() != ".xlsx":
            raise MappingError("微信群配对表必须使用 .xlsx 格式，请另存为 Excel 工作簿。")

    def create(self) -> Path:
        """Create an empty template, or validate an existing file without changing it."""
        with _WRITE_LOCK, self._process_write_lock():
            if self.path.exists():
                self.list_entries()
                return self.path
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = SHEET_NAME
            sheet.append(HEADERS)
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="23455B")
            for column, width in zip("ABCD", (38, 48, 32, 24)):
                sheet.column_dimensions[column].width = width
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = "A1:D1"
            try:
                self._save(workbook, expected=None)
            finally:
                workbook.close()
            return self.path

    def load(self) -> list[GroupMappingEntry]:
        return self.list_entries()

    def list_entries(self) -> list[GroupMappingEntry]:
        workbook, sheet, _original = self._open()
        try:
            return [entry for _row, entry in self._entries(sheet)]
        finally:
            workbook.close()

    def resolve(self, company_name: str) -> str:
        """Return the exact group name, failing closed if there is no unique mapping."""
        company_name = _name(company_name, "公司名称")
        for entry in self.list_entries():
            if entry.company_name == company_name:
                return entry.group_name
        raise MappingError(
            f"Excel 配对表中没有公司「{company_name}」的微信群。"
            "请先在配对表中新增公司名称和微信群完整名称并保存；不会使用公司名搜索或改发其他群。"
        )

    def upsert(
        self, company_name: str, group_name: str, notes: str | None = None,
        *, expected_group: str | None | object = _UNSPECIFIED,
    ) -> GroupMappingEntry:
        """Add/update one company after the caller has verified the WeChat group.

        The workbook must already exist. A missing workbook is never implicitly
        created during sending. ``notes=None`` preserves the existing notes.
        ``expected_group`` supports a long-running verification: None means the
        company was absent when it began; a string is the group seen then. When
        supplied, a changed mapping is rejected in the same load used to edit.
        """
        company_name = _name(company_name, "公司名称")
        group_name = _name(group_name, "微信群完整名称")
        if expected_group is not _UNSPECIFIED and expected_group is not None:
            expected_group = _name(expected_group, "原配对微信群完整名称")
        if notes is not None and (
            not isinstance(notes, str) or len(notes) > 32767
            or any(ord(char) < 32 and char not in "\t\n\r" for char in notes)
        ):
            raise MappingError("备注必须是有效的文本。")
        with _WRITE_LOCK, self._process_write_lock():
            self._check_excel_lock()
            workbook, sheet, original = self._open(for_write=True)
            try:
                entries = self._entries(sheet)
                if sheet.protection.sheet:
                    raise MappingError("配对工作表已受保护，请先在 Excel 中取消保护并保存。")
                existing = next(
                    (entry for _number, entry in entries if entry.company_name == company_name),
                    None,
                )
                actual_group = existing.group_name if existing is not None else None
                if expected_group is not _UNSPECIFIED and actual_group != expected_group:
                    raise MappingError(
                        f"公司「{company_name}」的 Excel 配对在等待或验证期间已发生变化。"
                        "请刷新配对并重新核对后再登记；现有配对未覆盖。"
                    )
                row = next(
                    (number for number, entry in entries if entry.company_name == company_name),
                    None,
                )
                if row is None:
                    row = self._first_empty_row(sheet)
                columns = self._optional_columns(sheet)
                if notes is not None and "备注" not in columns:
                    columns["备注"] = self._add_optional_column(sheet, "备注")
                if "更新时间" not in columns:
                    columns["更新时间"] = self._add_optional_column(sheet, "更新时间")
                self._set_text(sheet.cell(row, 1), company_name)
                self._set_text(sheet.cell(row, 2), group_name)
                if notes is not None:
                    self._set_text(sheet.cell(row, columns["备注"]), notes)
                updated_at = datetime.now().astimezone().isoformat(timespec="seconds")
                self._set_text(sheet.cell(row, columns["更新时间"]), updated_at)
                self._extend_filters(sheet, row)
                self._save(workbook, expected=original)
                saved_notes = (
                    _optional_text(sheet.cell(row, columns["备注"]).value, "备注")
                    if "备注" in columns else ""
                )
                return GroupMappingEntry(company_name, group_name, saved_notes, updated_at)
            finally:
                workbook.close()

    @staticmethod
    def _set_text(cell, value: str) -> None:
        cell.value = value
        # Even a note beginning '=' remains text; names additionally reject it.
        cell.data_type = "s"
        cell.number_format = "@"

    def _open(self, *, for_write: bool = False):
        try:
            original = self.path.read_bytes()
        except FileNotFoundError as exc:
            raise MappingError(
                f"尚未找到微信群配对表：{self.path}。"
                "请先点击“新建 Excel”或选择已有 Excel 表，填写公司名称和微信群完整名称并保存。"
            ) from exc
        except OSError as exc:
            raise MappingError("无法读取微信群配对表，请关闭 Excel 后检查文件路径和读取权限。") from exc
        if for_write:
            self._check_preservable(original)
        workbook = None
        try:
            with warnings.catch_warnings(record=True) as notices:
                warnings.simplefilter("always", UserWarning)
                workbook = load_workbook(BytesIO(original), data_only=False, rich_text=True)
            if for_write and notices:
                raise MappingError("此 Excel 包含无法完整保留的格式或功能，请把配对数据复制到新建配对表后重试。")
            matching = [
                sheet for sheet in workbook.worksheets
                if sheet.cell(1, 1).value == HEADERS[0] and sheet.cell(1, 2).value == HEADERS[1]
            ]
            if len(matching) != 1:
                raise MappingError(
                    "Excel 必须恰好有一张配对工作表，A1 为“公司名称”、B1 为“微信群完整名称”。"
                    "当前表头不符合要求或有多张配对表，原文件未修改。"
                )
            sheet = matching[0]
            if sheet.merged_cells.ranges:
                raise MappingError("配对工作表不能使用合并单元格，请在 Excel 中取消合并并核对每行配对。")
            return workbook, sheet, original
        except MappingError:
            if workbook is not None:
                workbook.close()
            raise
        except Exception as exc:
            if workbook is not None:
                workbook.close()
            raise MappingError("无法打开微信群配对表，请确认它是未损坏、未加密的 .xlsx 工作簿。") from exc

    @staticmethod
    def _populated_cells(sheet):
        # openpyxl is pinned to 3.1.5. Its public row/column iterators materialize
        # every missing cell up to max_row/max_column, including formatting at
        # XFD1048576. Read a snapshot of the existing sparse cell storage here;
        # never mutate it. All writes below still use the public cell API.
        return tuple(cell for cell in tuple(sheet._cells.values()) if cell.value is not None)

    @staticmethod
    def _optional_columns(sheet) -> dict[str, int]:
        columns = {}
        for cell in GroupMappingStore._populated_cells(sheet):
            if cell.row == 1 and cell.value in HEADERS[2:]:
                if cell.value in columns:
                    raise MappingError(f"Excel 表头“{cell.value}”重复，请保留唯一列后重试。")
                columns[cell.value] = cell.column
        return columns

    @staticmethod
    def _add_optional_column(sheet, header: str) -> int:
        # Formatting alone can extend max_column to Excel's last column XFD.
        # Appending after max_column would then create an invalid XFE column.
        # Reuse a column only when all its cells are empty, preserving extra
        # columns whose header is blank but whose data is still meaningful.
        occupied = {cell.column for cell in GroupMappingStore._populated_cells(sheet)}
        for column in range(3, 16385):
            if column not in occupied:
                sheet.cell(1, column, header)
                return column
        raise MappingError(
            "配对工作表没有可用的空白列，无法添加备注或更新时间。"
            "请整理空白列或把配对数据复制到新建配对表；原文件未修改。"
        )

    def _entries(self, sheet) -> list[tuple[int, GroupMappingEntry]]:
        columns = self._optional_columns(sheet)
        entries = []
        seen = {}
        rows = sorted({
            cell.row for cell in self._populated_cells(sheet)
            if cell.row >= 2 and cell.column in (1, 2)
        })
        for row in rows:
            company_cell, group_cell = sheet.cell(row, 1), sheet.cell(row, 2)
            if company_cell.value is None and group_cell.value is None:
                continue
            for cell, label in ((company_cell, "公司名称"), (group_cell, "微信群完整名称")):
                if cell.data_type == "f":
                    raise MappingError(f"Excel 第 {row} 行{label}是公式，请改为完整名称的纯文本。")
            company = _name(company_cell.value, f"Excel 第 {row} 行公司名称")
            group = _name(group_cell.value, f"Excel 第 {row} 行微信群完整名称")
            if company in seen:
                raise MappingError(
                    f"公司「{company}」在 Excel 第 {seen[company]}、{row} 行重复。"
                    "请仅保留一条配对，不会自动选择其中一行。"
                )
            seen[company] = row
            values = {
                header: _optional_text(sheet.cell(row, column).value, header)
                for header, column in columns.items()
            }
            entries.append((row, GroupMappingEntry(company, group, values.get("备注", ""), values.get("更新时间", ""))))
        return entries

    @staticmethod
    def _first_empty_row(sheet) -> int:
        # Formatting can make max_row thousands of rows longer than the actual
        # table. Reuse the first entirely empty data row, retaining notes-only
        # rows and any other cells containing values outside the mapping columns.
        occupied = {cell.row for cell in GroupMappingStore._populated_cells(sheet)}
        row = 2
        while row in occupied:
            row += 1
        if row > 1048576:
            raise MappingError("配对工作表没有可用的空白行，请整理数据或新建配对表；原文件未修改。")
        return row

    @staticmethod
    def _check_preservable(original: bytes) -> None:
        # openpyxl cannot round-trip every OOXML feature (notably shapes).
        # Refuse to rewrite such workbooks instead of silently dropping content.
        try:
            with ZipFile(BytesIO(original)) as archive:
                unsupported = (
                    "xl/drawings/", "xl/activeX/", "xl/ctrlProps/", "xl/embeddings/",
                    "xl/pivot", "xl/queryTables/", "xl/slicer", "customXml/", "xl/vbaProject",
                )
                if any(name.startswith(unsupported) for name in archive.namelist()):
                    raise MappingError(
                        "此 Excel 含图形、宏或数据连接等高级内容，无法安全更新。"
                        "请把公司和群名复制到本软件新建的配对表；原文件未修改。"
                    )
        except BadZipFile as exc:
            raise MappingError("配对表不是有效的 .xlsx 工作簿，原文件未修改。") from exc

    def _check_excel_lock(self) -> None:
        if self.path.with_name("~$" + self.path.name).exists():
            raise MappingError("微信群配对表正在 Excel 中打开，请先保存并关闭 Excel，再保存配对。")

    @contextmanager
    def _process_write_lock(self):
        # The byte comparison in _save detects edits by Excel, but is not an
        # atomic compare-and-swap: two app processes could both pass it before
        # either replaces the workbook. Lock the entire read/edit/replace cycle.
        # Keep the lock file in the OS temp directory and never unlink it:
        # unlinking a locked file would let another process lock a new inode.
        identity = os.path.normcase(str(self.path.resolve()))
        lock_name = "geshui-excel-" + sha256(identity.encode("utf-8")).hexdigest() + ".lock"
        try:
            stream = (Path(tempfile.gettempdir()) / lock_name).open("a+b")
        except OSError as exc:
            raise MappingError("无法取得配对表写入锁，请检查临时文件夹写入权限后重试。") from exc
        with stream:
            try:
                if os.name == "nt":
                    import msvcrt
                    stream.seek(0, os.SEEK_END)
                    if stream.tell() == 0:
                        stream.write(b"\0")
                        stream.flush()
                    stream.seek(0)
                    msvcrt.locking(stream.fileno(), msvcrt.LK_NBLCK, 1)
                else:
                    import fcntl
                    fcntl.flock(stream.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
            except OSError as exc:
                raise MappingError(
                    "另一个任务正在保存此 Excel 配对表，请稍后刷新配对再重试；现有文件未覆盖。"
                ) from exc
            # Closing the handle also releases the lock after an exception or
            # process exit, so no stale-lock recovery can discard a live lock.
            yield

    def _save(self, workbook, *, expected: bytes | None) -> None:
        temporary = None
        try:
            self._check_excel_lock()
            self.path.parent.mkdir(parents=True, exist_ok=True)
            descriptor, filename = tempfile.mkstemp(
                prefix=f".{self.path.stem}-", suffix=".xlsx", dir=self.path.parent,
            )
            os.close(descriptor)
            temporary = Path(filename)
            workbook.save(temporary)
            with temporary.open("rb+") as stream:
                os.fsync(stream.fileno())
            self._check_excel_lock()
            if expected is None:
                # Atomic no-clobber creation, including another process racing us.
                os.link(temporary, self.path)
            else:
                current = self.path.read_bytes()
                if sha256(current).digest() != sha256(expected).digest():
                    raise MappingError("Excel 配对表刚被其他程序修改，请重新加载后再保存，避免覆盖新内容。")
                os.replace(temporary, self.path)
        except MappingError:
            raise
        except FileExistsError as exc:
            raise MappingError("配对表已由其他程序创建，请重新加载；现有文件未覆盖。") from exc
        except PermissionError as exc:
            raise MappingError("无法保存配对表，请保存并关闭 Excel，检查文件夹写入权限后重试；原文件未修改。") from exc
        except OSError as exc:
            raise MappingError("无法保存配对表，请检查磁盘空间、文件是否存在及写入权限后重试。") from exc
        finally:
            if temporary is not None:
                temporary.unlink(missing_ok=True)

    @staticmethod
    def _extend_filters(sheet, row: int) -> None:
        if sheet.auto_filter.ref:
            left, top, right, bottom = range_boundaries(sheet.auto_filter.ref)
            if left <= 1 and right >= 2 and top == 1 and row > bottom:
                sheet.auto_filter.ref = f"{get_column_letter(left)}{top}:{get_column_letter(right)}{row}"
        for table in sheet.tables.values():
            left, top, right, bottom = range_boundaries(table.ref)
            if left <= 1 and right >= 2 and top == 1 and row == bottom + 1:
                table.ref = f"{get_column_letter(left)}{top}:{get_column_letter(right)}{row}"
                if table.autoFilter is not None:
                    table.autoFilter.ref = table.ref
